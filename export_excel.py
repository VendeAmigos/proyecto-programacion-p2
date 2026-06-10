import sqlite3
import pandas as pd
import os
import io
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def exportar_excel(in_memory=False):
    db_path = os.path.join(os.path.dirname(__file__), 'database', 'database.db')
    
    if not os.path.exists(db_path):
        print(f"Error: No se encontró la base de datos en {db_path}")
        return None

    print("Conectando a la base de datos...")
    con = sqlite3.connect(db_path)
    
    # Obtener todas las tablas
    query_tablas = "SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%';"
    tablas = pd.read_sql_query(query_tablas, con)['name'].tolist()
    
    if in_memory:
        output = io.BytesIO()
    else:
        output = os.path.join(os.path.dirname(__file__), 'base_de_datos_exportada.xlsx')
        print(f"Exportando tablas a {output}...")
    
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        # Pestaña de relaciones/metadatos
        info_relaciones = [
            {"Tabla": "productos", "Relacionado_con": "pedido_items (1 a N)", "Descripcion": "Catálogo de productos disponibles"},
            {"Tabla": "usuarios", "Relacionado_con": "pedidos, visitas, avisos, banners (1 a N)", "Descripcion": "Usuarios registrados y administradores"},
            {"Tabla": "pedidos", "Relacionado_con": "pedido_items (1 a N), usuarios (N a 1)", "Descripcion": "Órdenes de compra generales"},
            {"Tabla": "pedido_items", "Relacionado_con": "pedidos (N a 1), productos (N a 1)", "Descripcion": "Detalle de qué producto y cantidad va en cada pedido"},
            {"Tabla": "visitas", "Relacionado_con": "usuarios (N a 1)", "Descripcion": "Registro de rutas visitadas"},
            {"Tabla": "avisos", "Relacionado_con": "usuarios (N a 1)", "Descripcion": "Avisos globales del sistema"},
            {"Tabla": "banners", "Relacionado_con": "usuarios (N a 1)", "Descripcion": "Banners publicitarios de la portada"}
        ]
        df_info = pd.DataFrame(info_relaciones)
        df_info.to_excel(writer, sheet_name='Info_Relaciones', index=False)
        
        # Pestañas para cada tabla
        for tabla in tablas:
            df = pd.read_sql_query(f"SELECT * FROM {tabla}", con)
            df.to_excel(writer, sheet_name=tabla, index=False)
            print(f" - Tabla '{tabla}' exportada con {len(df)} registros.")

        # Aplicar diseño
        workbook = writer.book
        
        # Estilos
        header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
        
        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            
            # Ajustar anchos de columna y aplicar bordes
            for col in sheet.columns:
                max_length = 0
                column = col[0].column_letter # Get the column name
                for cell in col:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(cell.value)
                    except:
                        pass
                    # Aplicar bordes
                    cell.border = thin_border
                
                adjusted_width = (max_length + 2)
                sheet.column_dimensions[column].width = adjusted_width
            
            # Aplicar estilo al encabezado
            for cell in sheet[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
            
    con.close()
    
    if in_memory:
        output.seek(0)
        return output
    else:
        print("\n¡Exportación completada exitosamente!")
        print(f"Puedes abrir el archivo: {output}")

if __name__ == '__main__':
    exportar_excel()
